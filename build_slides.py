"""Build ISS_Tomato_Multi-Omics.pptx, an editable summary deck, from the
repository's figures. The numbers in the slide text are those of the
manuscript and of rnaseq/RERUN_2026-10-14.md and microbiome/RERUN_2026-10-15.md;
each slide's notes name the files they come from.

This deck replaces a NotebookLM-generated deck whose slides were flattened
images with claims the analyses do not support.
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, 'figures')
OUT = os.path.join(HERE, 'ISS_Tomato_Multi-Omics.pptx')

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x2F, 0x59, 0x85)
PAPER = RGBColor(0xFA, 0xFA, 0xF7)
W, H = Inches(13.333), Inches(7.5)
FOOTER = 'VEG-05 integrated omics · github.com/dr-richard-barker/veg05-integrated-omics'

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def text_box(slide, x, y, w, h, paragraphs, size=18, color=INK, bold=False):
    """paragraphs: list of str, or (str, level) for indented sub-points."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, p in enumerate(paragraphs):
        text, level = (p, 0) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = ('• ' if level == 0 and len(paragraphs) > 1 and not bold else '– ' if level else '') + text
        para.level = level
        para.space_after = Pt(8)
        for run in para.runs:
            run.font.size = Pt(size - 2 * level)
            run.font.color.rgb = color
            run.font.bold = bold
    return box


def add_image(slide, name, x, y, max_w, max_h):
    path = os.path.join(FIG, name)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, x + (max_w - w) // 2, y + (max_h - h) // 2, w, h)


def new_slide(title, notes, number=True):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    text_box(s, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0), [title], size=30, bold=True)
    if number:
        text_box(s, Inches(0.6), Inches(7.0), Inches(10), Inches(0.4), [FOOTER], size=10, color=MUTED)
        text_box(s, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.4), [str(len(prs.slides))], size=10, color=MUTED)
    s.notes_slide.notes_text_frame.text = notes
    return s


LEFT = (Inches(0.6), Inches(1.5), Inches(6.2), Inches(5.3))
RIGHT = (Inches(7.0), Inches(1.4), Inches(5.8), Inches(5.5))

# 1 -------------------------------------------------------------------------
s = new_slide('Integrated multi-omics of ISS VEG-05 tomatoes',
              'Sources: README.md; manuscript/latex/main.tex.', number=False)
text_box(s, Inches(0.6), Inches(1.4), Inches(12), Inches(1.2),
         ['Leaf and root microbiome (OSD-766) and host transcriptome (OSD-767) of tomato '
          'grown on the ISS under red-rich and blue-rich light'], size=20, color=MUTED)
add_image(s, 'fig1_study_design.png', Inches(1.4), Inches(2.6), Inches(10.5), Inches(3.9))
text_box(s, Inches(0.6), Inches(6.6), Inches(12), Inches(0.8),
         ['Data: NASA OSDR OSD-766 and OSD-767 (Spern et al. 2026; Dixit et al. 2026). '
          'Analyses re-run and checked in October 2026.'], size=12, color=MUTED)

# 2 -------------------------------------------------------------------------
s = new_slide('Study design',
              'Sources: manuscript/latex/main.tex (Methods); supplementary_tables/Table_S1_sample_metadata_rnaseq.csv; '
              'microbiome/track_stats_16S.tsv and track_stats_ITS.tsv (sample counts).')
text_box(s, *LEFT, [
    'Tomato cv. Red Robin grown on the ISS (VEG-05), with ground controls at Kennedy Space Center',
    'Red-rich (90% red : 10% blue) or blue-rich (50% red : 50% blue) LED light',
    'RNA-seq: 36 samples (21 leaf, 15 adventitious root)',
    'Amplicon sequencing: 16S from 142 samples and ITS from 141, across leaf, root, '
    'adventitious root, wick, soil and other compartments',
])
add_image(s, 'fig1_study_design.png', *RIGHT)

# 3 -------------------------------------------------------------------------
s = new_slide('Analysis pipeline',
              'Sources: README.md (pipeline table, software versions); microbiome/RERUN_2026-10-15.md; '
              'rnaseq/RERUN_2026-10-14.md; renv.lock.')
text_box(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.3), [
    'DADA2: 348 bacterial ASVs (after removing chloroplast and mitochondrial ASVs) and 78 fungal ASVs',
    'DESeq2 (flight × light model), WGCNA co-expression modules',
    'PERMANOVA, dysbiosis index and ALDEx2 for the microbiome',
    'MOFA+ and module–taxon correlation networks for integration',
    'FAPROTAX 1.2.12 for bacterial functions; genus-level FUNGuild categories for fungi',
    'Re-run in October 2026 (R 4.6.1, Bioconductor 3.23; renv.lock). The 16S results reproduce the '
    'original run exactly; ITS differs slightly (78 vs 77 ASVs) with newer DADA2 and UNITE versions',
])

# 4 -------------------------------------------------------------------------
s = new_slide('Spaceflight changes leaf bacterial and fungal communities',
              'Sources: microbiome/community_health/beta_diversity_permanova_{16S,ITS}.tsv (Table S4; terms added '
              'sequentially: flight, light, flight × light); alpha_diversity_16S.tsv (group means of Observed ASVs).')
text_box(s, *LEFT, [
    'PERMANOVA, flight term:',
    ('leaf bacteria R² 0.28 (p 0.001), leaf fungi R² 0.36 (p 0.007)', 1),
    ('adventitious-root bacteria R² 0.29 (p 0.001), fungi R² 0.40 (p 0.001)', 1),
    'Light had no significant main effect except on root fungi (R² 0.19, p 0.005)',
    'Flight × light interaction: bacteria in both tissues (R² 0.11) and root fungi (R² 0.16)',
    'Flight leaves had more bacterial ASVs: group means 17.8–25.8 vs 10.0–15.2 in ground leaves',
], size=17)
add_image(s, 'fig2a_alpha_diversity_16S.png', *RIGHT)

# 5 -------------------------------------------------------------------------
s = new_slide('Flight communities are displaced from ground communities',
              'Sources: microbiome/community_health/dysbiosis_index_{16S,ITS}.tsv (means ± SD of leaf samples). '
              'Ground indices average 1 by construction.')
text_box(s, Inches(0.6), Inches(1.5), Inches(12), Inches(1.9), [
    'Dysbiosis index: mean Bray–Curtis distance to ground samples ÷ mean distance among ground samples',
    'Leaf bacteria: flight 2.06 ± 0.99 vs ground 1.00 ± 0.47.  Leaf fungi: flight 3.69 ± 2.78 vs ground 1.00 ± 0.22',
], size=17)
add_image(s, 'fig2b_dysbiosis_16S.png', Inches(0.8), Inches(3.3), Inches(5.6), Inches(3.6))
add_image(s, 'fig2b_dysbiosis_ITS.png', Inches(6.9), Inches(3.3), Inches(5.6), Inches(3.6))

# 6 -------------------------------------------------------------------------
s = new_slide('Few taxa differ in abundance, all higher in flight',
              'Sources: microbiome/community_health/aldex2_results_*_flight.tsv (BH-adjusted Welch p < 0.05 and '
              '|effect| > 1; positive = higher in flight); taxonomy bootstrap support from the seeded DADA2 re-run.')
text_box(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.3), [
    'ALDEx2, BH-adjusted p < 0.05 and |effect| > 1:',
    ('Leaf: an Erwiniaceae ASV (ASV024, effect +1.10; Pantoea at 49% bootstrap support) and '
     'Penicillium citrinum (ASV003, +1.90)', 1),
    ('Adventitious root: Burkholderia-Caballeronia-Paraburkholderia (ASV004, +1.58), Rhizobiaceae '
     '(ASV021, +3.25), Dyadobacter (ASV022, +2.22)', 1),
    'A Fusarium ASV in roots (+1.91) narrowly missed the threshold (adjusted p 0.053)',
], size=18)

# 7 -------------------------------------------------------------------------
s = new_slide('The leaf transcriptional response to flight depends on light',
              'Sources: rnaseq/degs_summary.tsv (Table S2); DEGs: adjusted p < 0.05 and |log2FC| ≥ 1. Main effects '
              'are averaged over the other factor (rnaseq/RERUN_2026-10-14.md).')
text_box(s, *LEFT, [
    'Leaf: 4,716 DEGs under blue-rich light vs 523 under red-rich light (separate models); '
    '3,189 genes with a flight × light interaction',
    'Leaf flight main effect (averaged over light): 2,683 DEGs',
    'Adventitious root flight main effect: 1,501 DEGs, 83% up; interaction: 6 DEGs',
    'DEG: adjusted p < 0.05 and |log2 fold change| ≥ 1',
], size=17)
add_image(s, 'fig3_deg_summary.png', *RIGHT)

# 8 -------------------------------------------------------------------------
s = new_slide('WGCNA: one root module tracks fungal dysbiosis',
              'Sources: rnaseq/wgcna_{Leaf,Adv-Root}/module_trait_correlations.tsv (Pearson r; per-trait n); '
              'integration/networks/module_directionality.tsv (Table S5); microbiome/RERUN_2026-10-15.md.')
text_box(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(5.3), [
    'Leaf: 12 modules; turquoise (2,037 genes) correlates with flight (r = +0.73)',
    'Root: 16 modules; black (169 genes) correlates with ITS dysbiosis (r = −0.85, adjusted p 0.009, '
    '12 plants) but not with flight (r = −0.38, adjusted p 0.32)',
    'Caveat: the ITS dysbiosis index itself correlates with flight (r = 0.71); the black-module '
    'correlation holds within flight (r = −0.87) and within ground (r = −0.75) plants',
    'Classed as "microbe-driven"; its only enriched GO term is "nucleus"',
], size=16)
add_image(s, 'fig4_module_traits_AdvRoot.png', Inches(8.2), Inches(1.3), Inches(4.6), Inches(5.6))

# 9 -------------------------------------------------------------------------
s = new_slide('Host modules correlate with specific bacteria',
              'Sources: integration/networks/module_taxon_correlations.tsv (Spearman ρ, BH-adjusted; Table S9).')
text_box(s, *LEFT, [
    '31 significant module–taxon correlations (adjusted p < 0.05): 28 in leaf 16S, 3 in root ITS',
    'Flight-associated leaf turquoise module: Methylobacterium-Methylorubrum (ρ = +0.76), '
    'Burkholderia-Caballeronia-Paraburkholderia (+0.74), Azospirillum (+0.68)',
    'Roots: a Fusarium ASV with the blue (ρ = −0.93) and turquoise (ρ = +0.91) modules',
    'These are correlations; they do not show which side drives the other',
], size=17)
add_image(s, 'fig6_module_taxon_network.png', *RIGHT)

# 10 ------------------------------------------------------------------------
s = new_slide('MOFA+: a shared factor tracks flight, and light',
              'Sources: integration/factor_variance_explained.tsv; integration/factor_trait_correlations.tsv '
              '(Table S8; Spearman, BH-adjusted).')
text_box(s, *LEFT, [
    '18 leaf samples with RNA-seq, 16S and ITS from the same leaf',
    'Factor 1 explains 44% of transcriptome, 9% of 16S and 7% of ITS variance',
    'Factor 1 correlates with flight (ρ = −0.68, adjusted p 0.037) and with light (ρ = +0.63, '
    'adjusted p 0.053)',
    'Only one flight blue-light leaf has ITS data, so flight and light are partly confounded here',
], size=17)
add_image(s, 'fig5b_mofa_correlations.png', *RIGHT)

# 11 ------------------------------------------------------------------------
s = new_slide('Predicted functions',
              'Sources: microbiome/functional/faprotax_report_16S.tsv (Table S6; FAPROTAX 1.2.12, hyphens as word '
              'separators); microbiome/functional/funguild_summary_ITS.tsv (Table S7).')
text_box(s, *LEFT, [
    'FAPROTAX 1.2.12: 205 of 348 bacterial ASVs in 44 functional groups',
    'Largest by reads: chemoheterotrophy, ureolysis, nitrogen fixation, methylotrophy and methanol '
    'oxidation (mostly Methylobacterium-Methylorubrum); methanotrophy is minor',
    'Fungi (genus-level FUNGuild categories): 11 saprotroph ASVs (Penicillium, Aspergillus), '
    '2 plant-pathogen ASVs (Fusarium)',
    'Predicted from taxonomy, not measured',
], size=17)
add_image(s, 'fig7_faprotax.png', *RIGHT)

# 12 ------------------------------------------------------------------------
s = new_slide('Limitations',
              'Sources: manuscript/latex/main.tex (Limitations); microbiome/RERUN_2026-10-15.md (caveats).')
text_box(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.3), [
    'Small samples: 21 leaf and 15 root RNA-seq samples; 1–6 microbiome samples per flight × light group',
    'Leaf 16S: median 89 bacterial reads after removing plant reads; two flight root ITS samples kept one read each',
    'The ITS re-run differs slightly from the original run (newer DADA2 and UNITE versions)',
    'Three ground root samples were relabelled red-rich, resolving a conflict in the OSD-766 metadata',
    'All host–microbe associations are correlational',
], size=18)

# 13 ------------------------------------------------------------------------
s = new_slide('Code and data', 'Sources: README.md; LICENSE.')
text_box(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.3), [
    'Code: github.com/dr-richard-barker/veg05-integrated-omics (MIT licence)',
    'Site: dr-richard-barker.github.io/veg05-integrated-omics',
    'Data: NASA Open Science Data Repository, OSD-766 and OSD-767',
    'Re-run records: rnaseq/RERUN_2026-10-14.md and microbiome/RERUN_2026-10-15.md',
    'This deck is built from the repository by build_slides.py',
], size=18)

prs.save(OUT)
print(f'wrote {os.path.basename(OUT)}: {len(prs.slides)} slides')
